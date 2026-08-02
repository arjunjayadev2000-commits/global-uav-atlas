"""Export agent: CSV, Excel, SQLite, attribution, unresolved and coverage outputs.

Produces (all under ``output/``):

``uav_database.csv``          canonical records
``uav_source_records.csv``    the raw evidence layer
``photo_attribution.csv``     one row per atlas image with full licence chain
``unresolved_records.csv``    everything still open, with the precise blocker
``uav_database.xlsx``         8 sheets incl. the simple visual index
``uav_database.sqlite``       a consistent copy of the normalised database
``coverage_statistics.json``  machine readable coverage
``coverage_report.md``        the human readable version of the same numbers
"""

from __future__ import annotations

import csv
import json
import shutil
import sqlite3
from pathlib import Path
from typing import Any

from app.config import get_settings
from app.db import connect, query
from app.logging import AgentLogger, utc_now
from app.utils import country_iso3

LOG = AgentLogger("export_agent")
AGENT = "export_agent"


# ---------------------------------------------------------------------------
# Queries
# ---------------------------------------------------------------------------

CANONICAL_SQL = """
SELECT
    p.public_id                        AS record_id,
    p.canonical_name                   AS uav_name,
    COALESCE(c.name,'Unknown')         AS country_of_origin,
    COALESCE(c.iso3,'')                AS country_iso3,
    COALESCE(m.name,'')                AS manufacturer,
    COALESCE(f.name,'')                AS family,
    p.domain                           AS domain,
    COALESCE(p.category,'')            AS category,
    COALESCE(p.airframe_type,'')       AS airframe_type,
    COALESCE(p.status,'')              AS status,
    COALESCE(p.introduced_year,'')     AS introduced_year,
    COALESCE(p.retired_year,'')        AS retired_year,
    COALESCE(p.description_short,'')   AS description_short,
    p.verification_status              AS verification_status,
    p.confidence_score                 AS confidence_score,
    COALESCE(p.origin_confidence,'')   AS origin_confidence,
    COALESCE(p.participating_countries,'') AS participating_countries,
    COALESCE(p.production_country,'')  AS production_country,
    COALESCE(aliases.alias_list,'')    AS aliases,
    COALESCE(src.source_count,0)       AS source_count,
    COALESCE(src.source_list,'')       AS sources,
    COALESCE(img.atlas_path,'')        AS photo_path,
    COALESCE(img.source_page_url,'')   AS photo_source_page,
    COALESCE(img.license_name,'')      AS photo_license,
    COALESCE(img.identity_verdict,'no image') AS photo_identity_status,
    p.created_at                       AS created_at,
    p.updated_at                       AS updated_at
FROM platforms p
LEFT JOIN countries c     ON c.id = p.country_id
LEFT JOIN manufacturers m ON m.id = p.manufacturer_id
LEFT JOIN platform_families f ON f.id = p.family_id
LEFT JOIN images img      ON img.platform_id = p.id AND img.selected_for_atlas = 1
LEFT JOIN (
    SELECT platform_id, GROUP_CONCAT(alias, ' | ') AS alias_list
    FROM platform_aliases GROUP BY platform_id
) aliases ON aliases.platform_id = p.id
LEFT JOIN (
    SELECT ps.platform_id, COUNT(DISTINCT s.id) AS source_count,
           GROUP_CONCAT(DISTINCT s.title) AS source_list
    FROM platform_sources ps JOIN sources s ON s.id = ps.source_id
    GROUP BY ps.platform_id
) src ON src.platform_id = p.id
WHERE p.is_merged_into IS NULL
ORDER BY country_of_origin, manufacturer, uav_name
"""

ATTRIBUTION_SQL = """
SELECT
    p.public_id            AS record_id,
    p.canonical_name       AS uav_name,
    COALESCE(c.name,'Unknown') AS country_of_origin,
    i.local_path           AS local_path,
    i.atlas_path           AS atlas_path,
    i.source_url           AS image_url,
    i.source_page_url      AS source_page_url,
    COALESCE(i.photographer,'')     AS photographer,
    COALESCE(i.license_name,'')     AS license_name,
    COALESCE(i.license_url,'')      AS license_url,
    COALESCE(i.attribution_text,'') AS attribution_text,
    i.is_public_domain     AS public_domain,
    i.commercial_use       AS commercial_use_allowed,
    i.modification_allowed AS modification_allowed,
    i.license_verified     AS license_verified,
    i.width                AS width,
    i.height               AS height,
    i.mime_type            AS mime_type,
    i.sha256               AS sha256,
    i.perceptual_hash      AS perceptual_hash,
    i.identity_verdict     AS identity_verdict,
    i.identity_confidence  AS identity_confidence,
    COALESCE(i.identity_explanation,'') AS identity_explanation,
    i.retrieved_at         AS retrieved_at
FROM images i
JOIN platforms p ON p.id = i.platform_id
LEFT JOIN countries c ON c.id = p.country_id
WHERE i.selected_for_atlas = 1
ORDER BY uav_name
"""

UNRESOLVED_SQL = """
SELECT u.id, u.item_type, u.subject, u.detail, u.severity, u.suggested_action,
       COALESCE(u.blocker,'') AS blocker, u.status, u.entity_type, u.entity_id,
       u.created_at, u.updated_at
FROM unresolved_items u
WHERE u.status = 'open'
ORDER BY CASE u.severity WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END,
         u.item_type, u.subject
"""

SOURCE_RECORDS_SQL = """
SELECT rd.id AS raw_id, rd.raw_name, rd.source_dataset, rd.source_url, rd.external_ref,
       rd.raw_country, rd.raw_manufacturer, rd.domain, rd.category, rd.status,
       COALESCE(s.credibility_tier,3) AS credibility_tier,
       COALESCE(p.public_id,'') AS promoted_record_id,
       COALESCE(p.canonical_name,'') AS promoted_name,
       rd.created_at
FROM raw_discoveries rd
LEFT JOIN sources s ON s.id = rd.source_id
LEFT JOIN platforms p ON p.id = rd.promoted_platform_id
ORDER BY rd.source_dataset, rd.raw_name
"""

SOURCES_SQL = """
SELECT s.id, s.title, s.publisher, s.url, s.source_type, s.credibility_tier,
       s.retrieved_at, COALESCE(s.notes,'') AS notes,
       COUNT(ps.platform_id) AS platforms_supported
FROM sources s
LEFT JOIN platform_sources ps ON ps.source_id = s.id
GROUP BY s.id
ORDER BY s.credibility_tier, platforms_supported DESC
"""


def _rows(sql: str, conn: sqlite3.Connection | None = None) -> list[dict[str, Any]]:
    return [dict(r) for r in query(sql, conn=conn)]


def write_csv(path: Path, rows: list[dict[str, Any]], *, columns: list[str] | None = None) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = columns or (list(rows[0].keys()) if rows else [])
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, "") for c in cols})
    return len(rows)


# ---------------------------------------------------------------------------
# Coverage statistics
# ---------------------------------------------------------------------------


def coverage_statistics() -> dict[str, Any]:
    conn = connect()

    def one(sql: str, params: tuple[Any, ...] = ()) -> Any:
        rows = query(sql, params, conn)
        return rows[0][0] if rows else 0

    platforms = one("SELECT COUNT(*) FROM platforms WHERE is_merged_into IS NULL")
    merged = one("SELECT COUNT(*) FROM platforms WHERE is_merged_into IS NOT NULL")
    by_domain = {
        r["domain"]: r["n"]
        for r in query(
            "SELECT domain, COUNT(*) AS n FROM platforms WHERE is_merged_into IS NULL "
            "GROUP BY domain ORDER BY n DESC",
            conn=conn,
        )
    }
    by_country = {
        r["name"]: r["n"]
        for r in query(
            """
            SELECT COALESCE(c.name,'Unknown') AS name, COUNT(*) AS n
            FROM platforms p LEFT JOIN countries c ON c.id=p.country_id
            WHERE p.is_merged_into IS NULL GROUP BY name ORDER BY n DESC
            """,
            conn=conn,
        )
    }
    by_verification = {
        r["verification_status"]: r["n"]
        for r in query(
            "SELECT verification_status, COUNT(*) AS n FROM platforms "
            "WHERE is_merged_into IS NULL GROUP BY verification_status",
            conn=conn,
        )
    }
    by_category = {
        r["category"]: r["n"]
        for r in query(
            "SELECT COALESCE(NULLIF(category,''),'(unclassified)') AS category, COUNT(*) AS n "
            "FROM platforms WHERE is_merged_into IS NULL GROUP BY category "
            "ORDER BY n DESC LIMIT 40",
            conn=conn,
        )
    }
    real_countries = [
        name for name in by_country if name not in ("Unknown", "Multinational")
    ]

    images_selected = one("SELECT COUNT(*) FROM images WHERE selected_for_atlas=1")
    images_total = one("SELECT COUNT(*) FROM images")
    licensed = one(
        "SELECT COUNT(*) FROM images WHERE selected_for_atlas=1 AND license_verified=1"
    )
    identity_verified = one(
        "SELECT COUNT(*) FROM images WHERE selected_for_atlas=1 "
        "AND identity_verdict IN ('match','probable match')"
    )

    passes = [
        dict(r)
        for r in query(
            "SELECT pass_number, COUNT(*) AS strategies, SUM(candidates_seen) AS seen, "
            "SUM(new_candidates) AS new_candidates, SUM(CASE WHEN status='failed' THEN 1 ELSE 0 END) AS failed "
            "FROM discovery_runs GROUP BY pass_number ORDER BY pass_number",
            conn=conn,
        )
    ]

    unresolved_by_type = {
        r["item_type"]: r["n"]
        for r in query(
            "SELECT item_type, COUNT(*) AS n FROM unresolved_items WHERE status='open' "
            "GROUP BY item_type ORDER BY n DESC",
            conn=conn,
        )
    }

    stats: dict[str, Any] = {
        "generated_at": utc_now(),
        "canonical_platforms": platforms,
        "merged_duplicates": merged,
        "raw_source_records": one("SELECT COUNT(*) FROM raw_discoveries"),
        "manufacturers": one("SELECT COUNT(*) FROM manufacturers"),
        "countries_represented": len(real_countries),
        "countries_including_synthetic_buckets": len(by_country),
        "platform_families": one("SELECT COUNT(*) FROM platform_families"),
        "platform_variants": one("SELECT COUNT(*) FROM platform_variants"),
        "aliases": one("SELECT COUNT(*) FROM platform_aliases"),
        "sources": one("SELECT COUNT(*) FROM sources"),
        "platform_source_links": one("SELECT COUNT(*) FROM platform_sources"),
        "by_domain": by_domain,
        "by_verification_status": by_verification,
        "by_country": by_country,
        "by_category": by_category,
        "images": {
            "downloaded_total": images_total,
            "selected_for_atlas": images_selected,
            "licence_verified": licensed,
            "identity_match_or_probable": identity_verified,
            "platforms_with_image": images_selected,
            "platforms_without_image": platforms - images_selected,
            "coverage_pct": round(100.0 * images_selected / platforms, 2) if platforms else 0.0,
        },
        "discovery_passes": passes,
        "unresolved_open": one("SELECT COUNT(*) FROM unresolved_items WHERE status='open'"),
        "unresolved_by_type": unresolved_by_type,
        "possible_duplicates_open": one(
            "SELECT COUNT(*) FROM unresolved_items "
            "WHERE status='open' AND item_type='possible_duplicate'"
        ),
        "verification_events": one("SELECT COUNT(*) FROM verification_events"),
        "update_history_entries": one("SELECT COUNT(*) FROM update_history"),
    }
    stats["duplicate_review_rate_pct"] = (
        round(100.0 * stats["possible_duplicates_open"] / platforms, 2) if platforms else 0.0
    )
    return stats


def coverage_report_markdown(stats: dict[str, Any]) -> str:
    lines: list[str] = []
    add = lines.append
    add("# Global UAV Visual Atlas 2026 — Coverage Report")
    add("")
    add(f"Generated: **{stats['generated_at']}**")
    add("")
    add(
        "> Largest defensible consolidated registry generated from the configured sources. "
        "This is not a claim that every drone ever built is included."
    )
    add("")
    add("## Headline numbers")
    add("")
    add("| Metric | Value |")
    add("| --- | ---: |")
    add(f"| Canonical platforms | {stats['canonical_platforms']} |")
    add(f"| Raw source records | {stats['raw_source_records']} |")
    add(f"| Merged duplicates | {stats['merged_duplicates']} |")
    add(f"| Manufacturers / design organisations | {stats['manufacturers']} |")
    add(f"| Countries of origin represented | {stats['countries_represented']} |")
    add(f"| Platform families | {stats['platform_families']} |")
    add(f"| Recorded variants | {stats['platform_variants']} |")
    add(f"| Aliases | {stats['aliases']} |")
    add(f"| Sources | {stats['sources']} |")
    add(f"| Platform↔source links | {stats['platform_source_links']} |")
    add(f"| Atlas images | {stats['images']['selected_for_atlas']} |")
    add(f"| Image coverage | {stats['images']['coverage_pct']}% |")
    add(f"| Open unresolved items | {stats['unresolved_open']} |")
    add(f"| Duplicate review rate | {stats['duplicate_review_rate_pct']}% |")
    add("")

    add("## By domain")
    add("")
    add("| Domain | Platforms |")
    add("| --- | ---: |")
    for domain, count in stats["by_domain"].items():
        add(f"| {domain} | {count} |")
    add("")

    add("## By verification status")
    add("")
    add("| Status | Platforms |")
    add("| --- | ---: |")
    for status, count in stats["by_verification_status"].items():
        add(f"| {status} | {count} |")
    add("")

    add("## Country coverage")
    add("")
    add("| Country | ISO3 | Platforms |")
    add("| --- | --- | ---: |")
    for country, count in stats["by_country"].items():
        add(f"| {country} | {country_iso3(country) or '—'} | {count} |")
    add("")

    add("## Discovery passes")
    add("")
    add("| Pass | Strategies | Candidates seen | New candidates | Failed strategies |")
    add("| ---: | ---: | ---: | ---: | ---: |")
    for row in stats["discovery_passes"]:
        add(
            f"| {row['pass_number']} | {row['strategies']} | {row['seen'] or 0} | "
            f"{row['new_candidates'] or 0} | {row['failed'] or 0} |"
        )
    add("")

    add("## Image provenance")
    add("")
    images = stats["images"]
    add("| Metric | Value |")
    add("| --- | ---: |")
    add(f"| Images downloaded | {images['downloaded_total']} |")
    add(f"| Selected for the atlas | {images['selected_for_atlas']} |")
    add(f"| Licence verified | {images['licence_verified']} |")
    add(f"| Identity match or probable match | {images['identity_match_or_probable']} |")
    add(f"| Platforms without an image | {images['platforms_without_image']} |")
    add("")

    add("## Open unresolved items")
    add("")
    if stats["unresolved_by_type"]:
        add("| Type | Count |")
        add("| --- | ---: |")
        for item_type, count in stats["unresolved_by_type"].items():
            add(f"| {item_type} | {count} |")
    else:
        add("None.")
    add("")
    add("See `output/unresolved_records.csv` and `docs/UNRESOLVED.md` for the detail.")
    add("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

DATA_DICTIONARY = [
    ("platforms.public_id", "Stable atlas identifier derived from the normalised name."),
    ("platforms.canonical_name", "Preferred display name for the platform."),
    ("platforms.normalized_name", "Comparison key: accent folded, punctuation stripped, noise words removed."),
    ("platforms.domain", "Military | Civilian / Commercial | Dual-use | Research / Experimental | Unknown."),
    ("platforms.category", "Free-text role classification as reported by the strongest source."),
    ("platforms.status", "Programme/service status as reported by the source."),
    ("platforms.verification_status", "Verified | Probable | Needs Verification, recomputed from linked sources."),
    ("platforms.confidence_score", "0-1 score from source tier, source count, origin and manufacturer completeness."),
    ("platforms.origin_confidence", "High | Medium | Low | None - strength of the country-of-origin evidence."),
    ("platforms.participating_countries", "Partner nations for multinational programmes."),
    ("platforms.production_country", "Licensed-production country when it differs from design origin."),
    ("platforms.is_merged_into", "Set when this row was merged into another platform by the dedup agent."),
    ("countries.name", "Canonical country spelling (United States, United Kingdom, Türkiye, South Korea, ...)."),
    ("countries.iso3", "ISO 3166-1 alpha-3 code; null for Multinational and Unknown."),
    ("manufacturers.name", "Manufacturer or design organisation."),
    ("platform_aliases.alias_type", "alias | model_code | merged - provenance of the alternate name."),
    ("platform_variants.variant_role", "Why the variant exists, e.g. split from a combined designation."),
    ("images.local_path", "Original downloaded file on disk."),
    ("images.atlas_path", "Derivative embedded in the PDF/HTML atlas."),
    ("images.license_name", "Licence short name exactly as published by the host."),
    ("images.license_verified", "1 when the licence permits redistribution, commercial use and modification."),
    ("images.identity_verdict", "match | probable match | uncertain | mismatch | unverified."),
    ("images.identity_confidence", "0-1 confidence behind the identity verdict."),
    ("images.perceptual_hash", "Near-duplicate detection hash; prevents one photo serving two platforms."),
    ("images.selected_for_atlas", "Exactly one image per platform may carry this flag."),
    ("sources.credibility_tier", "1 official, 2 established institution, 3 open catalogue, 4 leads only."),
    ("raw_discoveries.fingerprint", "Deduplicating key over source + normalised name + variant + external ref."),
    ("unresolved_items.blocker", "Precise reason the item cannot be closed automatically."),
]


def write_excel(path: Path, data: dict[str, list[dict[str, Any]]], stats: dict[str, Any]) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        LOG.warning("openpyxl not installed - skipping Excel export")
        return False

    workbook = Workbook()

    def sheet(name: str, rows: list[dict[str, Any]], columns: list[str] | None = None) -> None:
        worksheet = workbook.create_sheet(title=name[:31])
        cols = columns or (list(rows[0].keys()) if rows else ["(no data)"])
        worksheet.append(cols)
        for row in rows:
            worksheet.append([row.get(c, "") for c in cols])
        header_fill = PatternFill("solid", fgColor="1F3864")
        for index, _ in enumerate(cols, start=1):
            cell = worksheet.cell(row=1, column=index)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            width = max(
                12,
                min(46, max([len(str(cols[index - 1]))] + [len(str(r.get(cols[index - 1], ""))) for r in rows[:200]] or [12]) + 2),
            )
            worksheet.column_dimensions[get_column_letter(index)].width = width
        worksheet.freeze_panes = "A2"
        if rows:
            worksheet.auto_filter.ref = worksheet.dimensions

    # Summary
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Global UAV Visual Atlas 2026 — Summary"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(["Generated", stats["generated_at"]])
    for label, key in (
        ("Canonical platforms", "canonical_platforms"),
        ("Raw source records", "raw_source_records"),
        ("Merged duplicates", "merged_duplicates"),
        ("Manufacturers", "manufacturers"),
        ("Countries represented", "countries_represented"),
        ("Platform families", "platform_families"),
        ("Variants", "platform_variants"),
        ("Sources", "sources"),
        ("Open unresolved items", "unresolved_open"),
        ("Duplicate review rate %", "duplicate_review_rate_pct"),
    ):
        summary.append([label, stats.get(key, "")])
    summary.append([])
    summary.append(["Images selected for atlas", stats["images"]["selected_for_atlas"]])
    summary.append(["Image coverage %", stats["images"]["coverage_pct"]])
    summary.append([])
    summary.append(["Domain", "Platforms"])
    for domain, count in stats["by_domain"].items():
        summary.append([domain, count])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 26

    sheet("UAV Database", data["canonical"])
    sheet(
        "Simple Visual Index",
        [
            {
                "Photo": r.get("photo_path", "") or "(no verified image)",
                "UAV Name": r["uav_name"],
                "Country of Origin": r["country_of_origin"],
            }
            for r in data["canonical"]
        ],
        ["Photo", "UAV Name", "Country of Origin"],
    )
    sheet("Source Records", data["raw"])
    sheet("Images", data["attribution"])
    sheet("Sources", data["sources"])
    sheet("Unresolved", data["unresolved"])
    sheet(
        "Data Dictionary",
        [{"Field": f, "Description": d} for f, d in DATA_DICTIONARY],
        ["Field", "Description"],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return True


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def export_sqlite(destination: Path) -> None:
    """Consistent copy of the live database via the SQLite backup API."""
    settings = get_settings()
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        destination.unlink()
    source = sqlite3.connect(str(settings.database_path))
    target = sqlite3.connect(str(destination))
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()


def run() -> dict[str, Any]:
    settings = get_settings()
    out = settings.paths.output
    out.mkdir(parents=True, exist_ok=True)
    conn = connect()

    data = {
        "canonical": _rows(CANONICAL_SQL, conn),
        "raw": _rows(SOURCE_RECORDS_SQL, conn),
        "attribution": _rows(ATTRIBUTION_SQL, conn),
        "unresolved": _rows(UNRESOLVED_SQL, conn),
        "sources": _rows(SOURCES_SQL, conn),
    }
    stats = coverage_statistics()

    written: dict[str, Any] = {}
    written["uav_database.csv"] = write_csv(out / "uav_database.csv", data["canonical"])
    written["uav_source_records.csv"] = write_csv(out / "uav_source_records.csv", data["raw"])
    written["photo_attribution.csv"] = write_csv(
        out / "photo_attribution.csv",
        data["attribution"],
        columns=[
            "record_id",
            "uav_name",
            "country_of_origin",
            "atlas_path",
            "image_url",
            "source_page_url",
            "photographer",
            "license_name",
            "license_url",
            "attribution_text",
            "public_domain",
            "commercial_use_allowed",
            "modification_allowed",
            "license_verified",
            "width",
            "height",
            "mime_type",
            "sha256",
            "perceptual_hash",
            "identity_verdict",
            "identity_confidence",
            "identity_explanation",
            "retrieved_at",
        ],
    )
    written["unresolved_records.csv"] = write_csv(out / "unresolved_records.csv", data["unresolved"])
    written["sources.csv"] = write_csv(out / "sources.csv", data["sources"])

    (out / "coverage_statistics.json").write_text(
        json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    (out / "coverage_report.md").write_text(coverage_report_markdown(stats), encoding="utf-8")
    written["coverage_statistics.json"] = True
    written["coverage_report.md"] = True

    written["uav_database.xlsx"] = write_excel(out / "uav_database.xlsx", data, stats)

    export_sqlite(out / "uav_database.sqlite")
    written["uav_database.sqlite"] = True

    # Keep the human-facing unresolved doc in sync with the database.
    _write_unresolved_doc(data["unresolved"], stats)

    LOG.info("exports written: %s", {k: v for k, v in written.items()})
    LOG.event("exports_complete", **{k: str(v) for k, v in written.items()})
    return {"written": written, "stats": stats}


def _write_unresolved_doc(rows: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    settings = get_settings()
    path = settings.paths.docs / "UNRESOLVED.md"
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row["item_type"], []).append(row)

    lines = [
        "# Unresolved items",
        "",
        f"Generated: {stats['generated_at']}",
        "",
        f"Open items: **{len(rows)}**. The machine-readable copy is "
        "`output/unresolved_records.csv`; this file is the narrative version.",
        "",
    ]
    for item_type, items in sorted(grouped.items(), key=lambda kv: -len(kv[1])):
        lines.append(f"## {item_type} ({len(items)})")
        lines.append("")
        blockers = {i["blocker"] for i in items if i.get("blocker")}
        if blockers:
            lines.append("Blockers observed: " + "; ".join(sorted(blockers)))
            lines.append("")
        actions = {i["suggested_action"] for i in items if i.get("suggested_action")}
        if actions:
            lines.append("Suggested action(s):")
            for action in sorted(actions)[:4]:
                lines.append(f"- {action}")
            lines.append("")
        lines.append("<details><summary>Affected subjects (first 60)</summary>")
        lines.append("")
        for item in items[:60]:
            detail = (item.get("detail") or "").replace("\n", " ")[:160]
            lines.append(f"- **{item['subject']}** — {detail}")
        if len(items) > 60:
            lines.append(f"- … and {len(items) - 60} more (see the CSV)")
        lines.append("")
        lines.append("</details>")
        lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")


def copy_seed_snapshot() -> None:
    """Keep an immutable copy of the seed alongside the exports for provenance."""
    settings = get_settings()
    seed = settings.paths.seed / "Global_UAV_Database_2026.sqlite"
    if seed.is_file():
        shutil.copy2(seed, settings.paths.output / "seed_snapshot.sqlite")
